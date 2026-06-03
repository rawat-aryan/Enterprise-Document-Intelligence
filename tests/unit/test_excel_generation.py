from __future__ import annotations

import io

import openpyxl

from backend.app.services.excel_service import ExcelService


def make_test_invoices(n: int = 5) -> list[dict]:
    return [
        {
            "vendor_name": f"Vendor {i}",
            "invoice_number": f"INV-{i:03d}",
            "invoice_date": "2024-01-15",
            "total_amount": 1000.0 * i,
            "tax_amount": 180.0 * i,
            "currency": "INR",
            "validation_status": "valid",
            "is_duplicate": i % 3 == 0,
            "duplicate_score": 0.9 if i % 3 == 0 else None,
            "confidence_score": 0.95,
            "validation_errors": [],
        }
        for i in range(1, n + 1)
    ]


def test_generate_invoice_report_returns_bytes():
    service = ExcelService()
    invoices = make_test_invoices(10)
    result = service.generate_invoice_report(invoices)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_generated_excel_has_correct_sheets():
    service = ExcelService()
    invoices = make_test_invoices(5)
    result = service.generate_invoice_report(invoices)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    sheet_names = wb.sheetnames
    assert "Invoice Summary" in sheet_names
    assert "Vendor Analytics" in sheet_names
    assert "Duplicate Detection" in sheet_names
    assert "Tax Analysis" in sheet_names
    assert "Exceptions" in sheet_names


def test_invoice_summary_has_data():
    service = ExcelService()
    invoices = make_test_invoices(3)
    result = service.generate_invoice_report(invoices)
    wb = openpyxl.load_workbook(io.BytesIO(result))
    ws = wb["Invoice Summary"]
    # Header row + 3 data rows
    assert ws.max_row >= 4


def test_empty_invoices():
    service = ExcelService()
    result = service.generate_invoice_report([])
    assert isinstance(result, bytes)
