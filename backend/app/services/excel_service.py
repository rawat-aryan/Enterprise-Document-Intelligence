from __future__ import annotations

import io
import logging
from typing import Any

logger = logging.getLogger(__name__)


class ExcelService:
    # Color palette
    HEADER_FILL = "1F4E79"
    SUBHEADER_FILL = "2E75B6"
    DUPLICATE_FILL = "FFE0E0"
    WARNING_FILL = "FFF3CD"
    SUCCESS_FILL = "D4EDDA"
    ACCENT_FILL = "E8F4FD"

    def generate_invoice_report(
        self,
        invoices: list[dict[str, Any]],
        vendor_analytics: list[dict[str, Any]] | None = None,
        duplicates: list[dict[str, Any]] | None = None,
    ) -> bytes:
        """Generate a 5-sheet Excel workbook with full styling."""
        try:
            from openpyxl import Workbook

            wb = Workbook()
            wb.remove(wb.active)  # Remove default sheet

            self._create_summary_sheet(wb, invoices)
            self._create_vendor_sheet(wb, vendor_analytics or self._compute_vendor_analytics(invoices))
            self._create_duplicate_sheet(wb, duplicates or [i for i in invoices if i.get("is_duplicate")])
            self._create_tax_sheet(wb, invoices)
            self._create_exceptions_sheet(wb, invoices)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.read()

        except Exception as e:
            logger.error(f"Excel generation failed: {e}")
            raise

    def _apply_header_style(self, ws, row_num: int, col_count: int, text: str, fill_color: str):
        from openpyxl.styles import Alignment, Font, PatternFill

        cell = ws.cell(row=row_num, column=1, value=text)
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col_count > 1:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=col_count)

    def _apply_column_header(self, ws, row_num: int, headers: list[str]):
        from openpyxl.styles import Alignment, Font, PatternFill

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(
                start_color=self.SUBHEADER_FILL,
                end_color=self.SUBHEADER_FILL,
                fill_type="solid",
            )
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _auto_width(self, ws):
        from openpyxl.utils import get_column_letter

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    def _create_summary_sheet(self, wb, invoices: list[dict]):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        ws = wb.create_sheet("Invoice Summary")
        ws.freeze_panes = "A3"

        headers = [
            "Invoice #",
            "Vendor Name",
            "Invoice Date",
            "Due Date",
            "Subtotal",
            "Tax Amount",
            "Total Amount",
            "Currency",
            "GST Number",
            "Status",
            "Confidence",
            "Duplicate",
        ]
        self._apply_header_style(ws, 1, len(headers), "Invoice Summary Report", self.HEADER_FILL)
        self._apply_column_header(ws, 2, headers)

        thin = Side(style="thin", color="DDDDDD")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row_idx, inv in enumerate(invoices, 3):
            is_dup = inv.get("is_duplicate", False)
            fill_color = self.DUPLICATE_FILL if is_dup else ("FFFFFF" if row_idx % 2 == 0 else self.ACCENT_FILL)
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            row_data = [
                inv.get("invoice_number", ""),
                inv.get("vendor_name", ""),
                inv.get("invoice_date", ""),
                inv.get("due_date", ""),
                inv.get("subtotal"),
                inv.get("tax_amount"),
                inv.get("total_amount"),
                inv.get("currency", "INR"),
                inv.get("gst_number", ""),
                inv.get("validation_status", ""),
                inv.get("confidence_score"),
                "YES" if is_dup else "NO",
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                cell.border = border
                if col in (5, 6, 7):  # Amount columns
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col == 11:  # Confidence
                    if val is not None:
                        cell.number_format = "0.00%"

        # Add totals row
        total_row = len(invoices) + 3
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        total_amounts = sum(i.get("total_amount") or 0 for i in invoices)
        total_tax = sum(i.get("tax_amount") or 0 for i in invoices)
        ws.cell(row=total_row, column=7, value=total_amounts).font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=total_tax).font = Font(bold=True)
        ws.cell(row=total_row, column=7).number_format = "#,##0.00"
        ws.cell(row=total_row, column=6).number_format = "#,##0.00"

        self._auto_width(ws)

    def _create_vendor_sheet(self, wb, vendor_data: list[dict]):
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import PatternFill

        ws = wb.create_sheet("Vendor Analytics")
        ws.freeze_panes = "A3"

        headers = ["Vendor Name", "Invoice Count", "Total Amount", "Average Amount", "Duplicate Rate (%)", "Risk Score"]
        self._apply_header_style(ws, 1, len(headers), "Vendor Analytics", self.HEADER_FILL)
        self._apply_column_header(ws, 2, headers)

        for row_idx, vendor in enumerate(vendor_data, 3):
            row_data = [
                vendor.get("vendor_name", ""),
                vendor.get("invoice_count", 0),
                vendor.get("total_amount", 0),
                vendor.get("average_amount", 0),
                vendor.get("duplicate_rate", 0) * 100,
                vendor.get("risk_score", 0),
            ]
            fill_color = self.WARNING_FILL if vendor.get("risk_score", 0) > 0.7 else "FFFFFF"
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                if col in (3, 4):
                    cell.number_format = "#,##0.00"

        # Add bar chart
        if len(vendor_data) > 0:
            chart = BarChart()
            chart.title = "Total Amount by Vendor"
            chart.y_axis.title = "Amount"
            chart.x_axis.title = "Vendor"
            chart.style = 10
            data_end = len(vendor_data) + 2
            data = Reference(ws, min_col=3, min_row=2, max_row=data_end)
            cats = Reference(ws, min_col=1, min_row=3, max_row=data_end)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.shape = 4
            ws.add_chart(chart, "H3")

        self._auto_width(ws)

    def _create_duplicate_sheet(self, wb, duplicates: list[dict]):
        from openpyxl.styles import Font, PatternFill

        ws = wb.create_sheet("Duplicate Detection")

        headers = [
            "Invoice #",
            "Vendor Name",
            "Invoice Date",
            "Total Amount",
            "Duplicate Score",
            "Duplicate Of",
            "Action Required",
        ]
        self._apply_header_style(ws, 1, len(headers), "Duplicate Invoice Detection", "C0392B")
        self._apply_column_header(ws, 2, headers)

        for row_idx, inv in enumerate(duplicates, 3):
            score = inv.get("duplicate_score", 0) or 0
            fill_color = "FFB3B3" if score >= 0.95 else self.WARNING_FILL
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            row_data = [
                inv.get("invoice_number", ""),
                inv.get("vendor_name", ""),
                inv.get("invoice_date", ""),
                inv.get("total_amount"),
                score,
                inv.get("duplicate_of", ""),
                "BLOCK" if score >= 0.95 else "REVIEW",
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                if col == 4:
                    cell.number_format = "#,##0.00"
                if col == 5:
                    cell.number_format = "0.00%"
                if col == 7:
                    cell.font = Font(bold=True, color="FF0000" if val == "BLOCK" else "FF6600")

        ws.cell(row=1, column=1).value = f"Duplicate Detection Report - {len(duplicates)} potential duplicates found"
        self._auto_width(ws)

    def _create_tax_sheet(self, wb, invoices: list[dict]):
        from openpyxl.styles import Font, PatternFill

        ws = wb.create_sheet("Tax Analysis")
        ws.freeze_panes = "A3"

        headers = ["Vendor Name", "Invoice #", "Subtotal", "CGST", "SGST", "IGST", "Total Tax", "Tax Rate (%)"]
        self._apply_header_style(ws, 1, len(headers), "Tax Analysis Report", "1A5276")
        self._apply_column_header(ws, 2, headers)

        total_cgst = total_sgst = total_igst = 0.0
        for row_idx, inv in enumerate(invoices, 3):
            tax_bd = inv.get("tax_breakdown") or {}
            cgst = tax_bd.get("cgst", 0) or 0
            sgst = tax_bd.get("sgst", 0) or 0
            igst = tax_bd.get("igst", 0) or 0
            subtotal = inv.get("subtotal") or 0
            total_tax = cgst + sgst + igst or inv.get("tax_amount") or 0
            tax_rate = (total_tax / subtotal * 100) if subtotal > 0 else 0

            total_cgst += cgst
            total_sgst += sgst
            total_igst += igst

            row_data = [
                inv.get("vendor_name", ""),
                inv.get("invoice_number", ""),
                subtotal,
                cgst,
                sgst,
                igst,
                total_tax,
                round(tax_rate, 2),
            ]
            fill_color = self.ACCENT_FILL if row_idx % 2 == 0 else "FFFFFF"
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                if col in (3, 4, 5, 6, 7):
                    cell.number_format = "#,##0.00"

        # Totals
        total_row = len(invoices) + 3
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        for col, val in [(4, total_cgst), (5, total_sgst), (6, total_igst), (7, total_cgst + total_sgst + total_igst)]:
            cell = ws.cell(row=total_row, column=col, value=val)
            cell.font = Font(bold=True)
            cell.number_format = "#,##0.00"

        self._auto_width(ws)

    def _create_exceptions_sheet(self, wb, invoices: list[dict]):
        from openpyxl.styles import Font, PatternFill

        ws = wb.create_sheet("Exceptions")

        headers = ["Invoice #", "Vendor", "Amount", "Status", "Issue Type", "Details", "Recommended Action"]
        self._apply_header_style(ws, 1, len(headers), "Processing Exceptions & Alerts", "922B21")
        self._apply_column_header(ws, 2, headers)

        exceptions = []
        for inv in invoices:
            errors = inv.get("validation_errors") or []
            for err in errors:
                issue_type = "WARNING" if err.startswith("WARNING") else "ERROR"
                exceptions.append(
                    {
                        "invoice_number": inv.get("invoice_number", ""),
                        "vendor_name": inv.get("vendor_name", ""),
                        "total_amount": inv.get("total_amount"),
                        "status": inv.get("validation_status", ""),
                        "issue_type": issue_type,
                        "detail": err,
                        "action": "Review and correct" if issue_type == "ERROR" else "Verify amount",
                    }
                )
            if inv.get("is_duplicate"):
                exceptions.append(
                    {
                        "invoice_number": inv.get("invoice_number", ""),
                        "vendor_name": inv.get("vendor_name", ""),
                        "total_amount": inv.get("total_amount"),
                        "status": "DUPLICATE",
                        "issue_type": "DUPLICATE",
                        "detail": f"Possible duplicate of invoice {inv.get('duplicate_of', 'unknown')}",
                        "action": "Block payment pending investigation",
                    }
                )

        for row_idx, exc in enumerate(exceptions, 3):
            issue = exc["issue_type"]
            fill_color = "FFB3B3" if issue == "ERROR" else (self.WARNING_FILL if issue == "WARNING" else "E8D5F5")
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            row_data = [
                exc["invoice_number"],
                exc["vendor_name"],
                exc["total_amount"],
                exc["status"],
                exc["issue_type"],
                exc["detail"],
                exc["action"],
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill = fill
                if col == 3 and val is not None:
                    cell.number_format = "#,##0.00"
                if col == 5:
                    cell.font = Font(bold=True)

        if not exceptions:
            ws.cell(row=3, column=1, value="No exceptions found - all invoices processed successfully")
            ws.cell(row=3, column=1).font = Font(color="2ECC71", bold=True)

        self._auto_width(ws)

    def _compute_vendor_analytics(self, invoices: list[dict]) -> list[dict]:
        from collections import defaultdict

        vendors: dict[str, dict] = defaultdict(
            lambda: {
                "invoice_count": 0,
                "total_amount": 0.0,
                "duplicate_count": 0,
            }
        )
        for inv in invoices:
            name = inv.get("vendor_name") or "Unknown"
            vendors[name]["invoice_count"] += 1
            vendors[name]["total_amount"] += inv.get("total_amount") or 0
            if inv.get("is_duplicate"):
                vendors[name]["duplicate_count"] += 1

        result = []
        for name, data in vendors.items():
            count = data["invoice_count"]
            result.append(
                {
                    "vendor_name": name,
                    "invoice_count": count,
                    "total_amount": data["total_amount"],
                    "average_amount": data["total_amount"] / count if count > 0 else 0,
                    "duplicate_rate": data["duplicate_count"] / count if count > 0 else 0,
                    "risk_score": min(data["duplicate_count"] / max(count, 1) + 0.1, 1.0),
                }
            )
        return sorted(result, key=lambda x: x["total_amount"], reverse=True)


excel_service = ExcelService()
